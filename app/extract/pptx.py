"""PowerPoint extraction.

Two halves, deliberately separated:

* **Images** come from LibreOffice (`soffice --headless --convert-to pdf`),
  the only faithful renderer available locally, handed off to `extract_pdf`.
  When soffice is absent we still emit one PNG per slide, painted from the
  slide's own text, so everything downstream can assume an image exists.
* **Text** always comes from `python-pptx`, which sees shape text *and*
  speaker notes and so is more faithful than the PDF text layer. We keep the
  PDF text only when python-pptx somehow returned less.
"""

from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

from app.contracts import SLIDE_MAX_PX, SLIDE_RENDER_SCALE, ExtractionError
from app.extract.pdf import extract_pdf, normalise_text, page_stem

__all__ = [
    "soffice_path",
    "rendering_available",
    "extract_pptx",
    "slide_texts",
    "NO_RENDERER_MARK",
]

NO_RENDERER_MARK = "[no renderer: text-only extraction]"

_SOFFICE_TIMEOUT = 180

_CANDIDATE_SOFFICE = (
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/opt/homebrew/bin/soffice",
    "/usr/local/bin/soffice",
    "/usr/bin/soffice",
)

_CANDIDATE_FONTS = (
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/System/Library/Fonts/Helvetica.ttc",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
)


def soffice_path() -> str | None:
    """Locate LibreOffice, or None. `LC_SOFFICE` overrides (tests use it)."""
    override = os.environ.get("LC_SOFFICE", "").strip()
    if override:
        return override if Path(override).exists() else None
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    for cand in _CANDIDATE_SOFFICE:
        if Path(cand).exists():
            return cand
    return None


def rendering_available() -> bool:
    """True when real slide rendering is possible (soffice present)."""
    return soffice_path() is not None


# --------------------------------------------------------------------------
# text layer (python-pptx)
# --------------------------------------------------------------------------


def _shape_text(shape) -> list[str]:
    """Recursively pull text out of one shape, including groups and tables."""
    out: list[str] = []
    try:
        if shape.shape_type is not None and str(shape.shape_type).startswith("GROUP"):
            for sub in shape.shapes:
                out.extend(_shape_text(sub))
            return out
    except Exception:  # noqa: BLE001 - some shapes have no usable shape_type
        pass
    try:
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                cells = [c.text.strip() for c in row.cells]
                line = " | ".join(c for c in cells if c)
                if line:
                    out.append(line)
            return out
    except Exception:  # noqa: BLE001
        pass
    try:
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                line = "".join(run.text for run in para.runs).strip()
                if not line:
                    line = (para.text or "").strip()
                if line:
                    out.append(line)
    except Exception:  # noqa: BLE001
        pass
    return out


def _notes_text(slide) -> str:
    try:
        if not slide.has_notes_slide:
            return ""
        frame = slide.notes_slide.notes_text_frame
        if frame is None:
            return ""
        return (frame.text or "").strip()
    except Exception:  # noqa: BLE001
        return ""


def slide_texts(path: str | Path) -> list[str]:
    """Per-slide text from python-pptx: all shape text, then speaker notes."""
    path = Path(path)
    try:
        prs = Presentation(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"cannot read pptx {path.name}: {exc}") from exc
    out: list[str] = []
    for slide in prs.slides:
        lines: list[str] = []
        for shape in slide.shapes:
            lines.extend(_shape_text(shape))
        body = "\n".join(lines).strip()
        notes = _notes_text(slide)
        if notes:
            body = (body + "\n\nSpeaker notes:\n" + notes).strip()
        out.append(normalise_text(body))
    return out


# --------------------------------------------------------------------------
# rendering
# --------------------------------------------------------------------------


def _convert_to_pdf(src: Path, tmp: Path) -> Path:
    """Run LibreOffice into an isolated profile so a running instance's
    profile lock cannot make this flaky."""
    exe = soffice_path()
    if exe is None:
        raise ExtractionError("soffice not available")
    profile = tmp / "loprofile"
    profile.mkdir(parents=True, exist_ok=True)
    outdir = tmp / "pdf"
    outdir.mkdir(parents=True, exist_ok=True)
    cmd = [
        exe,
        f"-env:UserInstallation=file://{profile}",
        "--headless",
        "--norestore",
        "--convert-to",
        "pdf",
        "--outdir",
        str(outdir),
        str(src),
    ]
    try:
        proc = subprocess.run(
            cmd, capture_output=True, text=True, timeout=_SOFFICE_TIMEOUT, check=False
        )
    except subprocess.TimeoutExpired as exc:
        raise ExtractionError(
            f"LibreOffice timed out after {_SOFFICE_TIMEOUT}s on {src.name}"
        ) from exc
    produced = sorted(outdir.glob("*.pdf"))
    if not produced:
        raise ExtractionError(
            f"LibreOffice produced no PDF for {src.name} "
            f"(exit {proc.returncode}): {(proc.stderr or proc.stdout or '').strip()[:400]}"
        )
    return produced[0]


def _font(size: int) -> ImageFont.ImageFont:
    for cand in _CANDIDATE_FONTS:
        if Path(cand).exists():
            try:
                return ImageFont.truetype(cand, size)
            except Exception:  # noqa: BLE001
                continue
    return ImageFont.load_default()


def _wrap(text: str, width_chars: int) -> list[str]:
    lines: list[str] = []
    for raw in text.split("\n"):
        if not raw.strip():
            lines.append("")
            continue
        words = raw.split()
        cur = ""
        for w in words:
            trial = f"{cur} {w}".strip()
            if len(trial) <= width_chars:
                cur = trial
            else:
                if cur:
                    lines.append(cur)
                cur = w
        if cur:
            lines.append(cur)
    return lines


def _paint_text_slide(text: str, dest: Path, max_px: int) -> None:
    """Fallback image: the slide's text on white, so downstream always has a
    PNG even with no renderer installed."""
    w = min(max_px, 1280)
    h = int(w * 9 / 16)
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    title_f = _font(34)
    body_f = _font(22)
    margin = 48
    y = margin
    lines = _wrap(text or "(no text)", 66)
    for i, line in enumerate(lines):
        f = title_f if i == 0 else body_f
        if y > h - margin:
            draw.text((margin, y), "...", fill="black", font=body_f)
            break
        draw.text((margin, y), line, fill="black", font=f)
        y += (44 if i == 0 else 30)
    img.save(dest, format="PNG")


# --------------------------------------------------------------------------
# entry point
# --------------------------------------------------------------------------


def extract_pptx(
    path: str | Path,
    out_dir: Path,
    *,
    scale: float = SLIDE_RENDER_SCALE,
    max_px: int = SLIDE_MAX_PX,
    **_kw,
) -> list[tuple[int, Path, str]]:
    path = Path(path)
    out_dir = Path(out_dir)
    if not path.exists():
        raise ExtractionError(f"no such file: {path}")
    out_dir.mkdir(parents=True, exist_ok=True)

    texts = slide_texts(path)          # authoritative text layer
    rendered: list[tuple[int, Path, str]] = []

    if rendering_available():
        with tempfile.TemporaryDirectory(prefix="lc-pptx-") as td:
            tmp = Path(td)
            pdf = _convert_to_pdf(path, tmp)
            rendered = extract_pdf(pdf, out_dir, scale=scale, max_px=max_px)
        if len(rendered) != len(texts):
            # Trust the render for images, python-pptx for text, match by index.
            print(
                f"[extract_pptx] slide-count mismatch for {path.name}: "
                f"{len(texts)} slides in pptx, {len(rendered)} rendered pages"
            )
        results: list[tuple[int, Path, str]] = []
        for i, png, pdf_text in rendered:
            pptx_text = texts[i] if i < len(texts) else ""
            best = pptx_text if len(pptx_text) >= len(pdf_text) else pdf_text
            (out_dir / f"{page_stem(i)}.txt").write_text(best, encoding="utf-8")
            results.append((i, png, best))
        return results

    # No renderer: paint the text and mark it plainly for the critic.
    results = []
    for i, text in enumerate(texts):
        stem = page_stem(i)
        png_path = out_dir / f"{stem}.png"
        marked = f"{NO_RENDERER_MARK}\n{text}".strip()
        _paint_text_slide(text, png_path, max_px)
        (out_dir / f"{stem}.txt").write_text(marked, encoding="utf-8")
        results.append((i, png_path, marked))
    return results
