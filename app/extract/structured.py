"""Deterministic, structure-preserving extraction for non-slide artifacts."""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import warnings as py_warnings
import zipfile
from collections import Counter
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

from app.contracts import ExtractionError, LearningObject, SourceLocator


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, (list, tuple)):
        return [_json_value(v) for v in value]
    if isinstance(value, dict):
        return {str(k): _json_value(v) for k, v in value.items()}
    return str(value)


def _object_id(artifact_id: str, object_type: str, locator: SourceLocator) -> str:
    seed = json.dumps(
        [artifact_id, object_type, locator.to_dict()], sort_keys=True, ensure_ascii=False
    ).encode("utf-8")
    return hashlib.sha256(seed).hexdigest()[:24]


def _obj(
    artifact_id: str,
    kind: str,
    object_type: str,
    *,
    text: str = "",
    data: dict[str, Any] | None = None,
    quality: float = 1.0,
    warnings: list[str] | None = None,
    **location: Any,
) -> LearningObject:
    locator = SourceLocator(artifact_id=artifact_id, kind=kind, **location)
    return LearningObject(
        id=_object_id(artifact_id, object_type, locator),
        artifact_id=artifact_id,
        object_type=object_type,
        locator=locator,
        text=text,
        data=_json_value(data or {}),
        source_version=artifact_id,
        quality=quality,
        warnings=warnings or [],
    )


def objects_from_pages(
    artifact_id: str,
    kind: str,
    pages: list[tuple[int, Path, str]],
    source_path: Path,
) -> list[LearningObject]:
    """Create exact page or slide objects and separate PowerPoint notes."""
    objects: list[LearningObject] = []
    is_pptx = kind == "pptx"
    notes: list[str] = []
    if is_pptx:
        try:
            from pptx import Presentation

            prs = Presentation(str(source_path))
            for slide in prs.slides:
                try:
                    frame = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
                    notes.append((frame.text or "").strip() if frame is not None else "")
                except Exception:  # noqa: BLE001
                    notes.append("")
        except Exception:  # noqa: BLE001
            notes = []
    for index, _image, text in pages:
        location = {"slide": index + 1} if is_pptx else {"page": index + 1}
        objects.append(_obj(
            artifact_id, kind, "slide" if is_pptx else "page", text=text,
            data={"index": index, "rendered_image": f"page-{index + 1:04d}.png"},
            **location,
        ))
        note = notes[index] if index < len(notes) else ""
        if note:
            objects.append(_obj(
                artifact_id, kind, "speaker_note", text=note,
                slide=index + 1, note=1,
            ))
    return objects


def extract_docx(path: Path, artifact_id: str) -> list[LearningObject]:
    try:
        from docx import Document

        doc = Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"cannot read docx {path.name}: {exc}") from exc
    objects: list[LearningObject] = []
    block = 0
    table_index = 0
    for item in doc.iter_inner_content():
        if hasattr(item, "rows"):
            table_index += 1
            rows = [[cell.text for cell in row.cells] for row in item.rows]
            objects.append(_obj(
                artifact_id, "docx", "table",
                text="\n".join(" | ".join(row) for row in rows),
                data={"rows": rows, "n_rows": len(rows), "n_columns": max((len(r) for r in rows), default=0)},
                block=block, fragment=f"table-{table_index}",
            ))
        else:
            text = (getattr(item, "text", "") or "").strip()
            if text:
                style = getattr(getattr(item, "style", None), "name", "") or ""
                prompt = text.endswith("?") or bool(re.match(
                    r"^(explain|describe|draw|identify|calculate|discuss|compare|write|complete)\b",
                    text, re.I,
                ))
                objects.append(_obj(
                    artifact_id, "docx", "prompt" if prompt else "document_block",
                    text=text, data={"style": style, "is_prompt": prompt}, block=block,
                ))
        block += 1
    image_index = 0
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            image_index += 1
            part = rel.target_part
            blob = part.blob
            objects.append(_obj(
                artifact_id, "docx", "image",
                data={
                    "relationship_id": rel.rId, "part_name": str(part.partname),
                    "content_type": str(part.content_type), "byte_length": len(blob),
                    "sha256": hashlib.sha256(blob).hexdigest(),
                },
                block=block, fragment=f"image-{image_index}",
            ))
            block += 1
    for rel in doc.part.rels.values():
        if "hyperlink" in rel.reltype and getattr(rel, "target_ref", ""):
            objects.append(_obj(
                artifact_id, "docx", "hyperlink", text=str(rel.target_ref),
                block=block, fragment=f"link-{block}",
            ))
            block += 1
    return objects


def _chart_title(chart: Any, fallback: str) -> str:
    try:
        if chart.title and chart.title.tx and chart.title.tx.rich:
            return " ".join(
                run.t for para in chart.title.tx.rich.p for run in para.r if getattr(run, "t", "")
            ) or fallback
    except Exception:  # noqa: BLE001
        pass
    return fallback


def extract_xlsx(path: Path, artifact_id: str) -> tuple[list[LearningObject], list[str]]:
    caught: list[Any] = []
    try:
        import openpyxl
        with py_warnings.catch_warnings(record=True) as records:
            py_warnings.simplefilter("always")
            workbook = openpyxl.load_workbook(path, data_only=False, read_only=False)
            cached = openpyxl.load_workbook(path, data_only=True, read_only=False)
            caught = list(records)
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"cannot read xlsx {path.name}: {exc}") from exc
    objects: list[LearningObject] = []
    warnings: list[str] = []
    for record in caught:
        message = str(record.message)
        if "Web Extension" in message:
            warning = (
                "Workbook web-extension semantics are unsupported; the original package and "
                "extension parts are preserved for exact recovery"
            )
        else:
            warning = f"Workbook parser warning: {type(record.message).__name__}"
        if warning not in warnings:
            warnings.append(warning)
    for sheet_index, sheet in enumerate(workbook.worksheets):
        cached_sheet = cached[sheet.title]
        merged = [str(rng) for rng in sheet.merged_cells.ranges]
        hidden_rows = [i for i, dim in sheet.row_dimensions.items() if dim.hidden]
        hidden_columns = [k for k, dim in sheet.column_dimensions.items() if dim.hidden]
        tables = sorted(str(name) for name in sheet.tables.keys())
        validations = [str(dv.sqref) for dv in sheet.data_validations.dataValidation]
        objects.append(_obj(
            artifact_id, "xlsx", "sheet", text=sheet.title,
            data={
                "index": sheet_index, "state": sheet.sheet_state,
                "dimensions": sheet.calculate_dimension(), "freeze_panes": str(sheet.freeze_panes or ""),
                "auto_filter": str(sheet.auto_filter.ref or ""), "merged_ranges": merged,
                "hidden_rows": hidden_rows, "hidden_columns": hidden_columns,
                "tables": tables, "conditional_formatting_rules": len(sheet.conditional_formatting),
                "data_validations": validations,
            },
            sheet=sheet.title, cell_range=sheet.calculate_dimension(),
        ))
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                cached_value = cached_sheet[cell.coordinate].value
                objects.append(_obj(
                    artifact_id, "xlsx", "cell", text=str(cell.value),
                    data={
                        "value": cell.value, "formula": cell.value if cell.data_type == "f" else None,
                        "cached_value": cached_value, "data_type": cell.data_type,
                        "number_format": cell.number_format, "style_id": cell.style_id,
                        "is_date": bool(cell.is_date),
                    },
                    sheet=sheet.title, cell_range=cell.coordinate,
                ))
        for table_name in sheet.tables.keys():
            table = sheet.tables[table_name]
            objects.append(_obj(
                artifact_id, "xlsx", "range", text=str(table_name),
                data={"name": str(table_name), "ref": table.ref, "style": str(table.tableStyleInfo.name if table.tableStyleInfo else "")},
                sheet=sheet.title, cell_range=table.ref,
            ))
        for chart_index, chart in enumerate(sheet._charts):  # openpyxl has no public chart collection
            title = _chart_title(chart, f"Chart {chart_index + 1}")
            series: list[dict[str, Any]] = []
            for item in getattr(chart, "ser", []):
                row: dict[str, Any] = {}
                for key in ("val", "cat", "xVal", "yVal"):
                    node = getattr(item, key, None)
                    ref = getattr(node, "numRef", None) or getattr(node, "strRef", None)
                    formula = getattr(ref, "f", None) if ref is not None else None
                    if formula:
                        row[key] = formula
                series.append(row)
            anchor = getattr(chart, "anchor", None)
            anchor_cell = ""
            try:
                anchor_cell = f"{anchor._from.col + 1},{anchor._from.row + 1}"
            except Exception:  # noqa: BLE001
                pass
            objects.append(_obj(
                artifact_id, "xlsx", "chart", text=title,
                data={"type": type(chart).__name__, "series": series, "anchor": anchor_cell},
                sheet=sheet.title, chart=title, fragment=f"chart-{chart_index + 1}",
            ))
    try:
        names = []
        for name, defined in workbook.defined_names.items():
            names.append({"name": name, "attr_text": defined.attr_text})
        if names:
            first_sheet = workbook.worksheets[0].title if workbook.worksheets else ""
            objects.append(_obj(
                artifact_id, "xlsx", "range", text="Workbook named ranges",
                data={"defined_names": names}, sheet=first_sheet, fragment="defined-names",
            ))
    except Exception as exc:  # noqa: BLE001
        warnings.append(f"defined names could not be read: {type(exc).__name__}")
    try:
        with zipfile.ZipFile(path) as package:
            for name in sorted(package.namelist()):
                lowered = name.lower()
                if "webextension" not in lowered and not lowered.startswith("customxml/"):
                    continue
                blob = package.read(name)
                objects.append(_obj(
                    artifact_id, "xlsx", "package_part", text=name,
                    data={"path": name, "byte_length": len(blob), "sha256": hashlib.sha256(blob).hexdigest()},
                    fragment=name,
                ))
    except (OSError, zipfile.BadZipFile) as exc:
        warnings.append(f"workbook package inventory unavailable: {type(exc).__name__}")
    return objects, warnings


def _decode_csv(data: bytes) -> tuple[str, str]:
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace"), "utf-8-replacement"


def _value_type(value: str) -> str:
    value = value.strip()
    if not value:
        return "missing"
    try:
        int(value)
        return "integer"
    except ValueError:
        pass
    try:
        float(value)
        return "number"
    except ValueError:
        pass
    if re.fullmatch(r"\d{4}[-/]\d{1,2}[-/]\d{1,2}.*", value):
        return "date_or_datetime"
    return "text"


def extract_csv(path: Path, artifact_id: str) -> list[LearningObject]:
    text, encoding = _decode_csv(path.read_bytes())
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(StringIO(text), dialect=dialect)
    fields = [str(f) for f in (reader.fieldnames or [])]
    stats = {field: {"missing": 0, "types": Counter(), "samples": [], "unique": set()} for field in fields}
    row_count = 0
    sample_rows: list[dict[str, str]] = []
    for row in reader:
        row_count += 1
        if len(sample_rows) < 20:
            sample_rows.append({field: str(row.get(field) or "") for field in fields})
        for field in fields:
            value = str(row.get(field) or "")
            stat = stats[field]
            kind = _value_type(value)
            stat["types"][kind] += 1
            if kind == "missing":
                stat["missing"] += 1
            elif len(stat["samples"]) < 5 and value not in stat["samples"]:
                stat["samples"].append(value)
            if len(stat["unique"]) <= 10000:
                stat["unique"].add(value)
    objects = [_obj(
        artifact_id, "csv", "dataset", text=path.name,
        data={
            "row_count": row_count, "column_count": len(fields), "encoding": encoding,
            "delimiter": dialect.delimiter, "fields": fields, "sample_rows": sample_rows,
        },
        fragment="dataset",
    )]
    for field in fields:
        stat = stats[field]
        non_missing_types = {k: v for k, v in stat["types"].items() if k != "missing"}
        inferred = max(non_missing_types, key=non_missing_types.get) if non_missing_types else "unknown"
        objects.append(_obj(
            artifact_id, "csv", "dataset_field", text=field,
            data={
                "inferred_type": inferred, "type_counts": dict(stat["types"]),
                "missing_count": stat["missing"], "unique_count": len(stat["unique"]),
                "samples": stat["samples"], "row_count": row_count,
            },
            dataset_field=field,
        ))
    return objects


def extract_ipynb(path: Path, artifact_id: str) -> list[LearningObject]:
    try:
        notebook = json.loads(path.read_text("utf-8"))
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"cannot read ipynb {path.name}: {exc}") from exc
    if not isinstance(notebook, dict) or not isinstance(notebook.get("cells"), list):
        raise ExtractionError(f"cannot read ipynb {path.name}: missing cells array")
    objects: list[LearningObject] = []
    for index, cell in enumerate(notebook["cells"]):
        if not isinstance(cell, dict):
            continue
        cell_type = str(cell.get("cell_type", "unknown"))
        source = cell.get("source", "")
        text = "".join(source) if isinstance(source, list) else str(source)
        objects.append(_obj(
            artifact_id, "ipynb", "notebook_cell", text=text,
            data={
                "cell_type": cell_type, "execution_count": cell.get("execution_count"),
                "metadata": cell.get("metadata", {}), "cell_id": cell.get("id", ""),
            },
            notebook_cell=index,
        ))
        for output_index, output in enumerate(cell.get("outputs", []) or []):
            if not isinstance(output, dict):
                continue
            output_text = output.get("text", output.get("data", {}).get("text/plain", ""))
            if isinstance(output_text, list):
                output_text = "".join(output_text)
            objects.append(_obj(
                artifact_id, "ipynb", "notebook_output", text=str(output_text or ""),
                data={
                    "output_type": output.get("output_type", ""),
                    "execution_count": output.get("execution_count"),
                    "ename": output.get("ename", ""), "evalue": output.get("evalue", ""),
                    "traceback": output.get("traceback", []),
                    "mime_types": sorted((output.get("data") or {}).keys()),
                    "data": output.get("data", {}), "name": output.get("name", ""),
                },
                notebook_cell=index, output=output_index,
            ))
    metadata = notebook.get("metadata", {})
    objects.append(_obj(
        artifact_id, "ipynb", "document_block", text="Notebook metadata",
        data={"metadata": metadata, "nbformat": notebook.get("nbformat"), "nbformat_minor": notebook.get("nbformat_minor")},
        fragment="notebook-metadata",
    ))
    return objects


def extract_structured(path: Path, artifact_id: str, kind: str) -> tuple[list[LearningObject], list[str]]:
    if kind == "docx":
        return extract_docx(path, artifact_id), []
    if kind == "xlsx":
        return extract_xlsx(path, artifact_id)
    if kind == "csv":
        return extract_csv(path, artifact_id), []
    if kind == "ipynb":
        return extract_ipynb(path, artifact_id), []
    raise ExtractionError(f"unsupported structured artifact: {kind}")
