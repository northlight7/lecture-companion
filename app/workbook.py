"""Deterministic workbook inspection and copy-only formula experiments."""

from __future__ import annotations

import ast
import hashlib
import json
import math
import re
import shutil
import statistics
import time
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import range_boundaries

from app.course import atomic_write_json

AUDIT_VERSION = 2
MAX_FORMULAS = 2_000
MAX_CHARTS = 200
MAX_REFERENCE_CELLS = 250_000


def _artifact(store, course_id: str, artifact_id: str):
    artifact = next((row for row in store.load_artifacts(course_id) if row.id == artifact_id), None)
    if artifact is None or artifact.kind != "xlsx" or artifact.extraction_status != "done":
        raise KeyError("No current workbook artifact was found in this course.")
    path = (store.course_dir(course_id) / artifact.stored_path).resolve()
    if store.course_dir(course_id) not in path.parents or not path.is_file():
        raise KeyError("The current workbook original is unavailable.")
    return artifact, path


def _audit_path(store, course_id: str, artifact_id: str) -> Path:
    return store.workbook_inspection_path(course_id, artifact_id)


def _formula_text(value: Any) -> str:
    if isinstance(value, str):
        return value
    return str(getattr(value, "text", value) or "")


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _split_reference(reference: str, current_sheet: str) -> tuple[str, str]:
    reference = reference.strip()
    if "!" not in reference:
        return current_sheet, reference.replace("$", "")
    sheet, coordinate = reference.rsplit("!", 1)
    if sheet.startswith("'") and sheet.endswith("'"):
        sheet = sheet[1:-1].replace("''", "'")
    return sheet, coordinate.replace("$", "")


def _table_reference(workbook, reference: str):
    match = re.fullmatch(r"([^\[]+)\[([^\]]+)\]", reference.strip())
    if not match:
        return None
    table_name, column_name = match.groups()
    for sheet in workbook.worksheets:
        for table in sheet.tables.values():
            if str(table.name) != table_name and str(table.displayName) != table_name:
                continue
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
            headers = [str(sheet.cell(min_row, col).value or "") for col in range(min_col, max_col + 1)]
            if column_name not in headers:
                return None
            col = min_col + headers.index(column_name)
            data_end = max_row - (1 if getattr(table, "totalsRowShown", False) else 0)
            return sheet.title, col, min_row + 1, data_end
    return None


def _resolved_values(cached, reference: str, current_sheet: str) -> dict[str, Any]:
    table = _table_reference(cached, reference)
    if table:
        sheet_name, min_col, min_row, max_row = table
        max_col = min_col
        coordinate = f"{cached[sheet_name].cell(min_row, min_col).coordinate}:{cached[sheet_name].cell(max_row, max_col).coordinate}"
    else:
        sheet_name, coordinate = _split_reference(reference, current_sheet)
        if sheet_name not in cached.sheetnames:
            raise ValueError(f"unknown sheet {sheet_name!r}")
        sheet = cached[sheet_name]
        try:
            min_col, min_row, max_col, max_row = range_boundaries(coordinate)
        except ValueError as exc:
            raise ValueError(f"unsupported reference {reference!r}") from exc
        min_row = min_row or 1
        max_row = max_row or sheet.max_row
        min_col = min_col or 1
        max_col = max_col or sheet.max_column
    if (max_row - min_row + 1) * (max_col - min_col + 1) > MAX_REFERENCE_CELLS:
        raise ValueError(f"reference {reference!r} exceeds the {MAX_REFERENCE_CELLS}-cell inspection limit")
    sheet = cached[sheet_name]
    cells = [
        sheet.cell(row, col)
        for row in range(min_row, max_row + 1)
        for col in range(min_col, max_col + 1)
    ]
    values = [cell.value for cell in cells]
    return {
        "reference": reference, "sheet": sheet_name, "range": coordinate,
        "values": values, "coordinates": [cell.coordinate for cell in cells],
    }


def _reference_evidence(resolved: dict[str, Any]) -> dict[str, Any]:
    values = resolved["values"]
    present = [value for value in values if value is not None]
    return {
        "reference": resolved["reference"],
        "locator": {"sheet": resolved["sheet"], "cell_range": resolved["range"]},
        "cell_count": len(values), "nonblank_count": len(present),
        "sample": [_json_value(value) for value in present[:8]],
    }


def _split_args(source: str) -> list[str]:
    args: list[str] = []
    start = 0
    depth = 0
    bracket = 0
    quoted = False
    index = 0
    while index < len(source):
        char = source[index]
        if char == '"':
            if quoted and index + 1 < len(source) and source[index + 1] == '"':
                index += 2
                continue
            quoted = not quoted
        elif not quoted:
            if char == "(": depth += 1
            elif char == ")": depth -= 1
            elif char == "[": bracket += 1
            elif char == "]": bracket -= 1
            elif char == "," and depth == 0 and bracket == 0:
                args.append(source[start:index].strip())
                start = index + 1
        index += 1
    args.append(source[start:].strip())
    return args


def _scalar(cached, token: str, current_sheet: str) -> Any:
    token = token.strip()
    if token.startswith('"') and token.endswith('"'):
        return token[1:-1].replace('""', '"')
    try:
        return float(token)
    except ValueError:
        resolved = _resolved_values(cached, token, current_sheet)
        if len(resolved["values"]) != 1:
            raise ValueError(f"{token!r} is not a scalar reference")
        return resolved["values"][0]


def _criterion(cached, source: str, current_sheet: str) -> Any:
    pieces = [part.strip() for part in source.split("&")]
    if len(pieces) > 1:
        return "".join(str(_scalar(cached, part, current_sheet)) for part in pieces)
    return _scalar(cached, source, current_sheet)


def _matches(value: Any, criterion: Any) -> bool:
    if not isinstance(criterion, str):
        return value == criterion
    match = re.match(r"^(<=|>=|<>|=|<|>)(.*)$", criterion)
    operator, expected = (match.group(1), match.group(2)) if match else ("=", criterion)
    try:
        left, right = float(value), float(expected)
    except (TypeError, ValueError):
        left, right = str(value or ""), expected
        if "*" in right or "?" in right:
            pattern = "^" + re.escape(right).replace(r"\*", ".*").replace(r"\?", ".") + "$"
            equal = re.match(pattern, left, re.IGNORECASE) is not None
            return not equal if operator == "<>" else equal
    return {
        "=": left == right, "<>": left != right, "<": left < right,
        ">": left > right, "<=": left <= right, ">=": left >= right,
    }[operator]


def _numeric(values: list[Any]) -> list[float]:
    return [float(value) for value in values if isinstance(value, (int, float)) and not isinstance(value, bool)]


def _evaluate_formula(cached, formula: str, current_sheet: str) -> tuple[str, Any, list[dict[str, Any]]]:
    formula = formula.strip()
    if not formula.startswith("="):
        return "unsupported", None, []
    body = formula[1:].strip()
    if body.startswith('"') and body.endswith('"'):
        return "verified", body[1:-1].replace('""', '"'), []
    function = re.fullmatch(r"([A-Za-z.]+)\((.*)\)", body, re.DOTALL)
    if function:
        name = function.group(1).upper()
        args = _split_args(function.group(2))
        supported = {"SUM", "AVERAGE", "MIN", "MAX", "MEDIAN", "COUNT", "COUNTA", "STDEV", "STDEV.S"}
        if name in supported:
            resolved = [_resolved_values(cached, arg, current_sheet) for arg in args]
            values = [value for item in resolved for value in item["values"]]
            numbers = _numeric(values)
            if name == "SUM": value = sum(numbers)
            elif name == "AVERAGE": value = statistics.mean(numbers)
            elif name == "MIN": value = min(numbers)
            elif name == "MAX": value = max(numbers)
            elif name == "MEDIAN": value = statistics.median(numbers)
            elif name == "COUNT": value = len(numbers)
            elif name == "COUNTA": value = sum(item is not None and item != "" for item in values)
            else: value = statistics.stdev(numbers)
            return "verified", value, [_reference_evidence(item) for item in resolved]
        if name in {"COUNTIF", "COUNTIFS", "SUMIFS"}:
            offset = 1 if name == "SUMIFS" else 0
            if name == "COUNTIF" and len(args) == 2:
                pairs = args
            else:
                pairs = args[offset:]
            if len(pairs) % 2:
                return "unsupported", None, []
            sum_range = _resolved_values(cached, args[0], current_sheet) if name == "SUMIFS" else None
            criteria_ranges = [_resolved_values(cached, pairs[index], current_sheet) for index in range(0, len(pairs), 2)]
            criteria = [_criterion(cached, pairs[index], current_sheet) for index in range(1, len(pairs), 2)]
            lengths = {len(item["values"]) for item in criteria_ranges}
            if sum_range: lengths.add(len(sum_range["values"]))
            if len(lengths) != 1:
                raise ValueError("criteria and sum ranges have different sizes")
            mask = [
                all(_matches(item["values"][row], criterion) for item, criterion in zip(criteria_ranges, criteria))
                for row in range(next(iter(lengths)))
            ]
            if name == "SUMIFS":
                value = sum(float(item) for item, keep in zip(sum_range["values"], mask) if keep and isinstance(item, (int, float)))
            else:
                value = sum(mask)
            resolved = ([sum_range] if sum_range else []) + criteria_ranges
            evidence = [_reference_evidence(item) for item in resolved]
            evidence.append({"criteria": [_json_value(value) for value in criteria], "matching_rows": sum(mask)})
            return "verified", value, evidence
        return "unsupported", None, []

    tokens = [token for token in Tokenizer(formula).items if token.type != "WHITE-SPACE"]
    expression: list[str] = []
    resolved: list[dict[str, Any]] = []
    values: dict[str, Any] = {}
    for token in tokens:
        if token.type == "OPERAND" and token.subtype == "RANGE":
            item = _resolved_values(cached, token.value, current_sheet)
            if len(item["values"]) != 1:
                return "unsupported", None, [_reference_evidence(item)]
            key = f"v{len(values)}"
            values[key] = item["values"][0]
            expression.append(key)
            resolved.append(item)
        elif token.type == "OPERAND" and token.subtype == "NUMBER":
            expression.append(token.value)
        elif token.type == "OPERATOR-INFIX" and token.value in {"+", "-", "*", "/", "^"}:
            expression.append("**" if token.value == "^" else token.value)
        elif token.type == "PAREN":
            expression.append(token.value)
        else:
            return "unsupported", None, [_reference_evidence(item) for item in resolved]
    tree = ast.parse("".join(expression), mode="eval")
    allowed = (ast.Expression, ast.BinOp, ast.UnaryOp, ast.Add, ast.Sub, ast.Mult, ast.Div, ast.Pow, ast.USub, ast.UAdd, ast.Name, ast.Load, ast.Constant)
    if any(not isinstance(node, allowed) for node in ast.walk(tree)):
        return "unsupported", None, [_reference_evidence(item) for item in resolved]
    value = eval(compile(tree, "<workbook-formula>", "eval"), {"__builtins__": {}}, values)
    return "verified", value, [_reference_evidence(item) for item in resolved]


def _unit(workbook, sheet, cell, formula: str) -> tuple[str, str]:
    number_format = str(cell.number_format or "")
    evidence = f"number format {number_format!r}"
    haystack = formula.lower()
    labels: list[str] = []
    value = sheet.cell(max(1, cell.row - 1), cell.column).value
    if isinstance(value, str): labels.append(value)
    for col in range(max(1, cell.column - 3), cell.column):
        value = sheet.cell(cell.row, col).value
        if isinstance(value, str): labels.append(value)
    for token in Tokenizer(formula).items:
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        try:
            resolved = _resolved_values(workbook, token.value, sheet.title)
        except ValueError:
            continue
        reference_sheet = workbook[resolved["sheet"]]
        first_coordinate = resolved["coordinates"][0]
        first_cell = reference_sheet[first_coordinate]
        if first_cell.row > 1:
            header = reference_sheet.cell(first_cell.row - 1, first_cell.column).value
            if isinstance(header, str): labels.append(header)
        labels.extend(str(item) for item in resolved["values"][:3] if isinstance(item, str))
    haystack += " " + " ".join(labels).lower()
    if "%" in number_format or any(word in haystack for word in ("percent", "percentage", "rate", "ratio", "share")):
        return "percentage", evidence + " and nearby labels"
    if "hk$000" in haystack or "hk$'000" in haystack:
        return "HKD thousands", "formula or nearby source label contains HK$000"
    if re.match(r"=(COUNT|COUNTA|COUNTIF|COUNTIFS)\b", formula, re.IGNORECASE):
        return "count", "the source formula is a counting function"
    if re.search(r"\bminutes?\b|\(min\)", haystack):
        return "minutes", "a referenced source label identifies minutes"
    if any(mark in number_format for mark in ("$", "£", "€", "¥")) or any(word in haystack for word in ("revenue", "price", "wage", "income", "cost", "sales")):
        return "currency or monetary amount", evidence + " and nearby labels"
    if cell.is_date or any(word in haystack for word in ("date", "month", "year", "time")):
        return "date or time", evidence + " and nearby labels"
    if any(word in haystack for word in ("frequency", "number of")):
        return "count", "formula or nearby source label indicates a count"
    return "not identified", "no unit marker was found in the number format, formula, or nearby labels"


def _compare(calculated: Any, cached_value: Any) -> tuple[bool | None, float | None]:
    if cached_value is None:
        return None, None
    if isinstance(calculated, (int, float)) and isinstance(cached_value, (int, float)):
        difference = float(calculated) - float(cached_value)
        return math.isclose(float(calculated), float(cached_value), rel_tol=1e-9, abs_tol=1e-9), difference
    return calculated == cached_value, None


def _object_maps(store, course_id: str, artifact_id: str):
    objects = store.load_learning_objects(course_id, artifact_id)
    cells = {(obj.locator.sheet, obj.locator.cell_range): obj for obj in objects if obj.object_type == "cell"}
    charts = {(obj.locator.sheet, obj.locator.chart): obj for obj in objects if obj.object_type == "chart"}
    return cells, charts


def _attach_source_object(evidence: dict[str, Any], cell_objects: dict) -> None:
    locator = evidence.get("locator")
    if not isinstance(locator, dict):
        return
    first = str(locator.get("cell_range", "")).split(":", 1)[0].replace("$", "")
    obj = cell_objects.get((locator.get("sheet"), first))
    evidence["object_id"] = getattr(obj, "id", "")


def _uses_cached_formula_precedent(workbook, intermediate: list[dict[str, Any]]) -> bool:
    for item in intermediate:
        reference = item.get("reference")
        locator = item.get("locator")
        if not reference or not isinstance(locator, dict):
            continue
        try:
            resolved = _resolved_values(workbook, str(reference), str(locator.get("sheet", "")))
        except ValueError:
            continue
        sheet = workbook[resolved["sheet"]]
        if any(sheet[coordinate].data_type == "f" for coordinate in resolved["coordinates"]):
            return True
    return False


def _chart_references(series) -> dict[str, str]:
    out: dict[str, str] = {}
    for key in ("val", "cat", "xVal", "yVal"):
        node = getattr(series, key, None)
        ref = getattr(node, "numRef", None) or getattr(node, "strRef", None)
        formula = getattr(ref, "f", None) if ref is not None else None
        if formula:
            out[key] = formula
    return out


def _chart_title(chart: Any, fallback: str) -> str:
    try:
        return " ".join(run.t for para in chart.title.tx.rich.p for run in para.r if getattr(run, "t", "")) or fallback
    except Exception:
        return fallback


def _inspect_chart(cached, sheet, chart, chart_index: int, chart_obj, cell_objects: dict) -> dict[str, Any]:
    title = _chart_title(chart, f"Chart {chart_index + 1}")
    rows: list[dict[str, Any]] = []
    diagnostics: list[dict[str, str]] = []
    category_counts: list[int] = []
    for number, series in enumerate(chart.ser, start=1):
        refs = _chart_references(series)
        evidence = {}
        for key, reference in refs.items():
            try:
                resolved = _resolved_values(cached, reference, sheet.title)
                evidence[key] = _reference_evidence(resolved)
                _attach_source_object(evidence[key], cell_objects)
                if key in {"cat", "xVal"}: category_counts.append(len([value for value in resolved["values"] if value is not None]))
            except ValueError as exc:
                evidence[key] = {"reference": reference, "error": str(exc)}
                diagnostics.append({"severity": "error", "message": f"Series {number} has an unresolved {key} range: {exc}."})
        value_key = "val" if "val" in evidence else "yVal" if "yVal" in evidence else ""
        category_key = "cat" if "cat" in evidence else "xVal" if "xVal" in evidence else ""
        if value_key and category_key and "cell_count" in evidence[value_key] and "cell_count" in evidence[category_key] and evidence[value_key]["cell_count"] != evidence[category_key]["cell_count"]:
            diagnostics.append({"severity": "error", "message": f"Series {number} has different value and category range lengths."})
        rows.append({"series": number, "references": refs, "evidence": evidence})
    if not title or title.startswith("Chart "):
        diagnostics.append({"severity": "warning", "message": "The chart has no descriptive source title."})
    if category_counts and len(chart.ser) > min(category_counts) and type(chart).__name__ in {"LineChart", "BarChart"}:
        diagnostics.append({
            "severity": "warning",
            "message": f"The chart has {len(chart.ser)} series but only {min(category_counts)} category values. If categories are time periods, inspect Switch Row/Column orientation.",
        })
    if not diagnostics:
        diagnostics.append({"severity": "info", "message": "Every inspected series has resolvable source ranges with aligned value and category lengths."})
    return {
        "object_id": getattr(chart_obj, "id", ""),
        "locator": {"sheet": sheet.title, "chart": title},
        "title": title, "chart_type": type(chart).__name__, "series": rows,
        "diagnostics": diagnostics,
        "provenance": "measured from the preserved workbook chart definition and current cached cells",
    }


def _formatting_and_filters(workbook) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        for group in sheet.conditional_formatting:
            for rule in group.rules:
                rows.append({
                    "kind": "conditional_formatting", "locator": {"sheet": sheet.title, "cell_range": str(group.sqref)},
                    "rule_type": str(rule.type), "operator": str(rule.operator or ""),
                    "formula": [_json_value(value) for value in (rule.formula or [])],
                    "priority": int(rule.priority or 0), "stop_if_true": bool(rule.stopIfTrue),
                    "provenance": "measured from the preserved workbook rule",
                })
        auto_filter = sheet.auto_filter
        if auto_filter.ref:
            row = {
                "kind": "auto_filter", "locator": {"sheet": sheet.title, "cell_range": str(auto_filter.ref)},
                "columns": [], "sort": [], "provenance": "measured from the preserved workbook filter definition",
            }
            for column in auto_filter.filterColumn:
                detail: dict[str, Any] = {"column_id": int(column.colId)}
                if column.filters:
                    detail["values"] = list(column.filters.filter)
                    detail["date_groups"] = [
                        {key: getattr(item, key) for key in ("year", "month", "day", "dateTimeGrouping") if getattr(item, key) is not None}
                        for item in column.filters.dateGroupItem
                    ]
                if column.customFilters:
                    detail["custom"] = [
                        {"operator": str(item.operator or "equal"), "value": _json_value(item.val)}
                        for item in column.customFilters.customFilter
                    ]
                    detail["and"] = bool(column.customFilters._and)
                if column.top10:
                    detail["top10"] = {"top": bool(column.top10.top), "value": column.top10.val, "filter_value": column.top10.filterVal}
                row["columns"].append(detail)
            if auto_filter.sortState:
                row["sort"] = [
                    {"range": item.ref, "descending": bool(item.descending)}
                    for item in auto_filter.sortState.sortCondition
                ]
            rows.append(row)
        if sheet.sheet_state != "visible" or any(dim.hidden for dim in sheet.row_dimensions.values()) or any(dim.hidden for dim in sheet.column_dimensions.values()):
            rows.append({
                "kind": "hidden_structure", "locator": {"sheet": sheet.title, "cell_range": sheet.calculate_dimension()},
                "sheet_state": sheet.sheet_state,
                "hidden_rows": [index for index, dim in sheet.row_dimensions.items() if dim.hidden],
                "hidden_columns": [index for index, dim in sheet.column_dimensions.items() if dim.hidden],
                "provenance": "measured from the preserved workbook visibility state",
            })
    return rows


def inspect_workbook(store, course_id: str, artifact_id: str, *, force: bool = False) -> dict[str, Any]:
    artifact, path = _artifact(store, course_id, artifact_id)
    audit_path = _audit_path(store, course_id, artifact_id)
    if audit_path.is_file() and not force:
        cached_audit = json.loads(audit_path.read_text("utf-8"))
        if cached_audit.get("audit_version") == AUDIT_VERSION and cached_audit.get("source_hash") == artifact.content_hash:
            return {**cached_audit, "cache_hit": True, "stale": False}
    source_before = hashlib.sha256(path.read_bytes()).hexdigest()
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=False)
    cached = openpyxl.load_workbook(path, data_only=True, read_only=False)
    cell_objects, chart_objects = _object_maps(store, course_id, artifact_id)
    formulas: list[dict[str, Any]] = []
    formula_total = 0
    for sheet in workbook.worksheets:
        cached_sheet = cached[sheet.title]
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type != "f":
                    continue
                formula_total += 1
                if len(formulas) >= MAX_FORMULAS:
                    continue
                formula = _formula_text(cell.value)
                cached_value = cached_sheet[cell.coordinate].value
                try:
                    state, calculated, intermediate = _evaluate_formula(cached, formula, sheet.title)
                    for item in intermediate:
                        _attach_source_object(item, cell_objects)
                    uses_cached_precedent = state == "verified" and _uses_cached_formula_precedent(workbook, intermediate)
                    if uses_cached_precedent:
                        state = "verified_using_cached_precedents"
                    matches, difference = _compare(calculated, cached_value) if state == "verified" else (None, None)
                    error = ""
                except Exception as exc:
                    state, calculated, intermediate, matches, difference = "error", None, [], None, None
                    error = f"{type(exc).__name__}: {exc}"
                unit, unit_basis = _unit(workbook, sheet, cell, formula)
                obj = cell_objects.get((sheet.title, cell.coordinate))
                formulas.append({
                    "object_id": getattr(obj, "id", ""),
                    "locator": {"sheet": sheet.title, "cell_range": cell.coordinate},
                    "formula": formula, "cached_value": _json_value(cached_value),
                    "calculated_value": _json_value(calculated), "verification_state": state,
                    "matches_cached": matches, "difference": difference, "error": error,
                    "unit": unit, "unit_basis": unit_basis, "intermediate": intermediate,
                    "calculation_basis": (
                        "independent raw source values" if state == "verified"
                        else "includes one or more cached formula precedents" if state == "verified_using_cached_precedents"
                        else "not independently evaluated"
                    ),
                    "provenance": "computed locally from exact referenced cells; cached value remains labeled source data",
                })
    charts: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        for index, chart in enumerate(sheet._charts):
            if len(charts) >= MAX_CHARTS:
                continue
            title = _chart_title(chart, f"Chart {index + 1}")
            charts.append(_inspect_chart(cached, sheet, chart, index, chart_objects.get((sheet.title, title)), cell_objects))
    formatting = _formatting_and_filters(workbook)
    for finding in formatting:
        _attach_source_object(finding, cell_objects)
    source_after = hashlib.sha256(path.read_bytes()).hexdigest()
    verified_count = sum(row["verification_state"] == "verified" for row in formulas)
    mismatch_count = sum(row["verification_state"] == "verified" and row["matches_cached"] is False for row in formulas)
    result = {
        "audit_version": AUDIT_VERSION, "artifact_id": artifact.id,
        "source_hash": artifact.content_hash, "artifact_version": artifact.version,
        "source_path": artifact.source_path, "generated_at": time.time(),
        "cache_hit": False, "stale": False, "remote_model_used": False,
        "source_bytes_preserved": source_before == source_after == artifact.content_hash,
        "summary": {
            "sheets": len(workbook.sheetnames), "formula_count": formula_total,
            "formulas_inspected": len(formulas), "formulas_independently_verified": verified_count,
            "cached_mismatches": mismatch_count, "charts": len(charts),
            "formatting_filter_findings": len(formatting),
        },
        "formulas": formulas, "charts": charts, "formatting_and_filters": formatting,
        "uncertainty": [
            "Unsupported formulas are labeled unsupported and are not presented as verified.",
            "Cached values come from the source workbook. Independent values use the documented local evaluator, not Excel.",
            "Unit labels are inferred only from source number formats, formulas, and nearby labels, with the basis shown.",
            "Chart orientation warnings are structural checks, not claims about the author's intent.",
        ],
    }
    atomic_write_json(audit_path, result)
    return result


def audit_status(store, course_id: str, artifact_id: str) -> dict[str, Any]:
    artifact, _path = _artifact(store, course_id, artifact_id)
    path = _audit_path(store, course_id, artifact_id)
    if not path.is_file():
        if artifact.supersedes:
            previous = _audit_path(store, course_id, artifact.supersedes)
            if previous.is_file():
                audit = json.loads(previous.read_text("utf-8"))
                return {
                    "exists": True, "stale": True, "artifact_id": artifact_id,
                    "previous_artifact_id": artifact.supersedes,
                    "source_hash": audit.get("source_hash"),
                    "audit_version": audit.get("audit_version"), "summary": audit.get("summary", {}),
                }
        return {"exists": False, "stale": False, "artifact_id": artifact_id}
    audit = json.loads(path.read_text("utf-8"))
    return {
        "exists": True, "stale": audit.get("source_hash") != artifact.content_hash,
        "artifact_id": artifact_id, "source_hash": audit.get("source_hash"),
        "audit_version": audit.get("audit_version"), "summary": audit.get("summary", {}),
    }


def formula_experiment(
    store, course_id: str, artifact_id: str, *, sheet_name: str, cell_coordinate: str, formula: str,
) -> dict[str, Any]:
    artifact, source = _artifact(store, course_id, artifact_id)
    if not formula.startswith("=") or len(formula) > 2_000:
        raise ValueError("A bounded Excel formula beginning with = is required.")
    workbook = openpyxl.load_workbook(source, data_only=False, read_only=False)
    cached = openpyxl.load_workbook(source, data_only=True, read_only=False)
    if sheet_name not in workbook.sheetnames or not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", cell_coordinate.upper()):
        raise ValueError("Choose an existing sheet and a valid cell coordinate.")
    coordinate = cell_coordinate.upper()
    raw = json.dumps([AUDIT_VERSION, artifact.content_hash, sheet_name, coordinate, formula], separators=(",", ":"))
    experiment_id = "experiment-" + hashlib.sha256(raw.encode("utf-8")).hexdigest()[:20]
    run_dir = store.workbook_experiment_dir(course_id, artifact_id, experiment_id)
    result_path = run_dir / "result.json"
    if result_path.is_file():
        return {**json.loads(result_path.read_text("utf-8")), "cache_hit": True}
    run_dir.mkdir(parents=True, exist_ok=True)
    copy_path = run_dir / "workbook-copy.xlsx"
    shutil.copy2(source, copy_path)
    workbook[sheet_name][coordinate] = formula
    workbook.save(copy_path)
    state, value, intermediate = _evaluate_formula(cached, formula, sheet_name)
    cell_objects, _chart_objects = _object_maps(store, course_id, artifact_id)
    for item in intermediate:
        _attach_source_object(item, cell_objects)
    if state == "verified" and _uses_cached_formula_precedent(workbook, intermediate):
        state = "verified_using_cached_precedents"
    obj = cell_objects.get((sheet_name, coordinate))
    unit, unit_basis = _unit(workbook, workbook[sheet_name], workbook[sheet_name][coordinate], formula)
    result = {
        "experiment_id": experiment_id, "artifact_id": artifact_id,
        "source_hash": artifact.content_hash, "artifact_version": artifact.version,
        "sheet": sheet_name, "cell": coordinate, "formula": formula,
        "verification_state": state, "calculated_value": _json_value(value),
        "intermediate": intermediate, "unit": unit, "unit_basis": unit_basis,
        "calculation_basis": (
            "independent raw source values" if state == "verified"
            else "includes one or more cached formula precedents" if state == "verified_using_cached_precedents"
            else "not independently evaluated"
        ),
        "object_id": getattr(obj, "id", ""), "cache_hit": False,
        "source_bytes_preserved": hashlib.sha256(source.read_bytes()).hexdigest() == artifact.content_hash,
        "provenance": "modeled formula experiment on a course-scoped copy; the calculated value is local and the source workbook is unchanged",
        "remote_model_used": False,
    }
    atomic_write_json(result_path, result)
    return result


__all__ = ["inspect_workbook", "audit_status", "formula_experiment"]
