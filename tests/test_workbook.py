"""Deterministic spreadsheet inspection and copy-only experiment gates."""

from __future__ import annotations

import hashlib
from io import BytesIO
import zipfile
from xml.etree import ElementTree

import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule

from app.importer import import_files
from app.workbook import audit_status, formula_experiment, inspect_workbook


def _workbook_bytes(*, total: int = 60) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    data = workbook.create_sheet("Data")
    data.append(["Region", "Revenue (HK$000)"])
    data.append(["A", 10])
    data.append(["B", 20])
    data.append(["C", 30])
    summary["A1"] = "Total revenue"
    summary["B1"] = "=SUM(Data!B2:B4)"
    summary["A2"] = "Average revenue"
    summary["B2"] = "=B1/3"
    summary["A4"] = 1
    summary["A5"] = 2
    summary["A6"] = 3
    summary.conditional_formatting.add("A4:A6", CellIsRule(operator="greaterThan", formula=["1"], stopIfTrue=True))
    summary.conditional_formatting.add("A4:A6", CellIsRule(operator="lessThan", formula=["3"], stopIfTrue=False))
    data.auto_filter.ref = "A1:B4"
    data.row_dimensions[4].hidden = True
    chart = BarChart()
    for column in range(3, 7):
        summary.cell(4, column, f"Series {column}")
        summary.cell(5, column, column)
        summary.cell(6, column, column + 1)
        chart.add_data(Reference(summary, min_col=column, min_row=4, max_row=6), titles_from_data=True)
    chart.set_categories(Reference(summary, min_col=1, min_row=5, max_row=6))
    chart.title = "Orientation check"
    summary.add_chart(chart, "H2")
    stream = BytesIO()
    workbook.save(stream)
    source = stream.getvalue()

    output = BytesIO()
    with zipfile.ZipFile(BytesIO(source)) as incoming, zipfile.ZipFile(output, "w") as outgoing:
        for item in incoming.infolist():
            data_bytes = incoming.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(data_bytes)
                namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                for cell in root.iter(namespace + "c"):
                    if cell.attrib.get("r") not in {"B1", "B2"}:
                        continue
                    value = cell.find(namespace + "v")
                    if value is None:
                        value = ElementTree.SubElement(cell, namespace + "v")
                    value.text = str(total if cell.attrib["r"] == "B1" else total / 3)
                data_bytes = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
            outgoing.writestr(item, data_bytes)
    return output.getvalue()


def _import(store, payload: bytes):
    course = store.create_course("Workbook checks")
    result = import_files(store, course.id, [("L2/checks.xlsx", payload)])
    return course, result.filed[0].artifact_id


def test_audit_independently_verifies_formulas_units_chart_and_rules_without_changing_source(store):
    payload = _workbook_bytes()
    course, artifact_id = _import(store, payload)
    audit = inspect_workbook(store, course.id, artifact_id)
    assert audit["source_bytes_preserved"] is True
    assert audit["remote_model_used"] is False
    assert audit["summary"]["formula_count"] == 2
    assert audit["summary"]["formulas_independently_verified"] == 1
    assert audit["summary"]["cached_mismatches"] == 0
    total = next(row for row in audit["formulas"] if row["locator"]["cell_range"] == "B1")
    assert total["formula"] == "=SUM(Data!B2:B4)"
    assert total["cached_value"] == 60
    assert total["calculated_value"] == 60
    assert total["matches_cached"] is True
    assert total["calculation_basis"] == "independent raw source values"
    assert total["unit"] == "HKD thousands"
    assert total["intermediate"][0]["locator"] == {"sheet": "Data", "cell_range": "B2:B4"}
    chart = audit["charts"][0]
    assert chart["object_id"]
    assert any("Switch Row/Column" in row["message"] for row in chart["diagnostics"])
    rules = [row for row in audit["formatting_and_filters"] if row["kind"] == "conditional_formatting"]
    assert len(rules) == 2 and rules[0]["priority"] == 1 and rules[0]["stop_if_true"] is True
    assert any(row["kind"] == "auto_filter" for row in audit["formatting_and_filters"])
    assert any(row["kind"] == "hidden_structure" for row in audit["formatting_and_filters"])
    artifact = store.load_artifacts(course.id)[0]
    original = store.course_dir(course.id) / artifact.stored_path
    assert hashlib.sha256(original.read_bytes()).hexdigest() == artifact.content_hash == hashlib.sha256(payload).hexdigest()


def test_audit_labels_mismatch_and_is_idempotent(store):
    course, artifact_id = _import(store, _workbook_bytes(total=99))
    first = inspect_workbook(store, course.id, artifact_id)
    assert first["summary"]["cached_mismatches"] == 1
    assert first["formulas"][0]["matches_cached"] is False
    second = inspect_workbook(store, course.id, artifact_id)
    assert second["cache_hit"] is True and second["generated_at"] == first["generated_at"]


def test_formula_experiment_uses_copy_and_exact_source_identity(store):
    payload = _workbook_bytes()
    course, artifact_id = _import(store, payload)
    result = formula_experiment(
        store, course.id, artifact_id,
        sheet_name="Summary", cell_coordinate="B2", formula="=Data!B2*2",
    )
    assert result["verification_state"] == "verified"
    assert result["calculated_value"] == 20
    assert result["object_id"]
    assert result["source_bytes_preserved"] is True
    assert result["provenance"].startswith("modeled formula experiment on a course-scoped copy")
    copied = store.workbook_experiment_dir(course.id, artifact_id, result["experiment_id"]) / "workbook-copy.xlsx"
    assert copied.is_file() and store.course_dir(course.id) in copied.resolve().parents
    again = formula_experiment(
        store, course.id, artifact_id,
        sheet_name="Summary", cell_coordinate="B2", formula="=Data!B2*2",
    )
    assert again["cache_hit"] is True and again["experiment_id"] == result["experiment_id"]


def test_workbook_inspection_is_course_scoped_and_replaced_sources_do_not_serve_old_audits(store):
    course, artifact_id = _import(store, _workbook_bytes())
    inspect_workbook(store, course.id, artifact_id)
    other = store.create_course("Foreign workbook")
    with pytest.raises(KeyError):
        inspect_workbook(store, other.id, artifact_id)
    replaced = import_files(store, course.id, [("L2/checks.xlsx", _workbook_bytes(total=61))])
    new_id = replaced.filed[0].artifact_id
    assert new_id != artifact_id
    with pytest.raises(KeyError):
        audit_status(store, course.id, artifact_id)
    status = audit_status(store, course.id, new_id)
    assert status["exists"] is True and status["stale"] is True
    assert status["previous_artifact_id"] == artifact_id
