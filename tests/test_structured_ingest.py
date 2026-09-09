"""Six-format import, exact locators, hierarchy, dedupe, and quarantine."""

from __future__ import annotations

import csv
import json
import zipfile
from io import BytesIO, StringIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from PIL import Image

from app.importer import import_files
from app.questions import evidence_text


def _docx() -> bytes:
    out = BytesIO()
    doc = Document()
    doc.add_heading("ER modelling tutorial", 1)
    doc.add_paragraph("Draw an entity relationship diagram?")
    picture = BytesIO()
    Image.new("RGB", (12, 8), "blue").save(picture, format="PNG")
    picture.seek(0)
    doc.add_picture(picture)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Entity"
    table.cell(0, 1).text = "Identifier"
    table.cell(1, 0).text = "Student"
    table.cell(1, 1).text = "student_id"
    doc.save(out)
    return out.getvalue()


def _xlsx() -> bytes:
    out = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Returns"
    ws.append(["week", "return", "double"])
    ws.append([1, 0.02, "=B2*2"])
    ws.append([2, -0.01, "=B3*2"])
    ws.auto_filter.ref = "A1:C3"
    ws.freeze_panes = "A2"
    ws.row_dimensions[3].hidden = True
    ws.conditional_formatting.add("B2:B3", CellIsRule(operator="lessThan", formula=["0"], stopIfTrue=True))
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart.title = "Weekly returns"
    ws.add_chart(chart, "E2")
    wb.save(out)
    return out.getvalue()


def _ipynb() -> bytes:
    return json.dumps({
        "cells": [
            {"cell_type": "markdown", "id": "intro", "metadata": {}, "source": ["# Regression\n"]},
            {"cell_type": "code", "id": "fit", "metadata": {}, "execution_count": 3,
             "source": ["print('rmse')"], "outputs": [
                 {"output_type": "stream", "name": "stdout", "text": ["rmse\n"]},
                 {"output_type": "error", "ename": "ValueError", "evalue": "bad metric", "traceback": ["trace"]},
                 {"output_type": "display_data", "data": {
                     "image/png": "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
                 }, "metadata": {}},
             ]},
        ],
        "metadata": {"kernelspec": {"name": "python3"}}, "nbformat": 4, "nbformat_minor": 5,
    }).encode()


def test_all_new_formats_import_with_typed_exact_locators(store):
    course = store.create_course("Mixed course")
    csv_data = b"date,price,return\n2026-01-01,100,\n2026-01-08,101,0.01\n"
    result = import_files(store, course.id, [
        ("L2/tutorial.docx", _docx()),
        ("L2/workbook.xlsx", _xlsx()),
        ("L2/prices.csv", csv_data),
        ("L2/analysis.ipynb", _ipynb()),
    ])
    assert result.errors == [], result.errors
    assert {row.kind for row in result.filed} == {"docx", "xlsx", "csv", "ipynb"}
    assert all(row.object_count > 0 and row.status == "done" for row in result.filed)
    artifacts = store.load_artifacts(course.id)
    assert {a.source_path for a in artifacts} == {
        "L2/tutorial.docx", "L2/workbook.xlsx", "L2/prices.csv", "L2/analysis.ipynb"
    }

    objects = store.load_learning_objects(course.id)
    kinds = {obj.object_type for obj in objects}
    assert {"prompt", "table", "sheet", "cell", "chart", "dataset", "dataset_field",
            "notebook_cell", "notebook_output", "formatting_rule", "filter_rule"} <= kinds
    assert all(obj.locator.artifact_id == obj.artifact_id for obj in objects)
    formula = next(obj for obj in objects if obj.object_type == "cell" and obj.locator.cell_range == "C2")
    assert formula.data["formula"] == "=B2*2"
    assert formula.locator.sheet == "Returns"
    assert '"formula": "=B2*2"' in evidence_text(formula)
    sheet = next(obj for obj in objects if obj.object_type == "sheet")
    assert '"auto_filter": "A1:C3"' in evidence_text(sheet)
    rule = next(obj for obj in objects if obj.object_type == "formatting_rule")
    assert rule.locator.cell_range == "B2:B3" and rule.data["stop_if_true"] is True
    filter_rule = next(obj for obj in objects if obj.object_type == "filter_rule")
    assert filter_rule.locator.cell_range == "A1:C3"
    output = next(obj for obj in objects if obj.object_type == "notebook_output" and obj.data["output_type"] == "error")
    assert output.locator.notebook_cell == 1 and output.locator.output == 1
    field = next(obj for obj in objects if obj.object_type == "dataset_field" and obj.text == "return")
    assert field.data["missing_count"] == 1 and field.locator.dataset_field == "return"
    assert '"missing_count": 1' in evidence_text(field)
    dataset = next(obj for obj in objects if obj.object_type == "dataset")
    assert '"column_count": 3' in evidence_text(dataset) and '"row_count": 2' in evidence_text(dataset)


def test_incremental_unchanged_import_does_not_duplicate_objects(store):
    course = store.create_course("Incremental")
    data = b"x,y\n1,2\n"
    first = import_files(store, course.id, [("Week 1/data.csv", data)])
    before = store.load_learning_objects(course.id)
    second = import_files(store, course.id, [("Week 1/data.csv", data)])
    after = store.load_learning_objects(course.id)
    assert first.errors == second.errors == []
    assert second.filed[0].status == "unchanged"
    assert [obj.id for obj in after] == [obj.id for obj in before]
    assert len(store.load_artifacts(course.id)) == 1


def test_identical_bytes_keep_two_hierarchy_references_and_one_raw_copy(store):
    course = store.create_course("Repeated data")
    data = b"date,value\n2026-01-01,1\n"
    result = import_files(store, course.id, [
        ("L1/shared.csv", data), ("L2/shared.csv", data),
    ])
    assert result.errors == []
    artifacts = store.load_artifacts(course.id)
    assert len(artifacts) == 2
    assert {a.source_path for a in artifacts} == {"L1/shared.csv", "L2/shared.csv"}
    duplicate = next(a for a in artifacts if a.duplicate_of)
    assert duplicate.duplicate_of in {a.id for a in artifacts}
    assert len(list(store.raw_dir(course.id).glob("*"))) == 1


def test_changed_path_creates_version_and_supersedes_old_objects(store):
    course = store.create_course("Versions")
    import_files(store, course.id, [("L1/data.csv", b"x\n1\n")])
    result = import_files(store, course.id, [("L1/data.csv", b"x\n1\n2\n")])
    assert result.errors == []
    current = store.load_artifacts(course.id)
    history = store.load_artifacts(course.id, include_superseded=True)
    assert len(current) == 1 and len(history) == 2
    assert current[0].version == 2 and current[0].supersedes
    assert current[0].content_hash != next(a for a in history if a.id == current[0].supersedes).content_hash


def test_malformed_artifact_is_quarantined_without_aborting_batch(store):
    course = store.create_course("Quarantine")
    result = import_files(store, course.id, [
        ("bad.xlsx", b"not a workbook"),
        ("good.csv", b"a,b\n1,2\n"),
    ])
    assert len(result.filed) == 1 and result.filed[0].filename == "good.csv"
    assert len(result.errors) == 1 and "bad.xlsx" in result.errors[0]
    artifacts = store.load_artifacts(course.id)
    failed = next(a for a in artifacts if a.filename == "bad.xlsx")
    assert failed.extraction_status == "failed" and failed.extraction_quality == 0.0
    assert failed.extraction_warnings


def test_original_bytes_round_trip_through_api(store):
    course = store.create_course("Originals")
    data = b"a,b\n1,2\n"
    result = import_files(store, course.id, [("L1/source.csv", data)])
    artifact_id = result.filed[0].artifact_id
    from app.main import app

    with TestClient(app) as client:
        response = client.get(f"/api/courses/{course.id}/artifacts/{artifact_id}/original")
    assert response.status_code == 200 and response.content == data


def test_typed_objects_are_structurally_course_isolated(store):
    course_a = store.create_course("Analytics A")
    course_b = store.create_course("Analytics B")
    a = import_files(store, course_a.id, [("data.csv", b"feature,value\nA_ONLY_SIGNAL,1\n")])
    b = import_files(store, course_b.id, [("data.csv", b"feature,value\nB_ONLY_SIGNAL,2\n")])
    objects_b = store.load_learning_objects(course_b.id)
    encoded_b = json.dumps([obj.to_dict() for obj in objects_b])
    assert "B_ONLY_SIGNAL" in encoded_b and "A_ONLY_SIGNAL" not in encoded_b
    assert store.load_learning_objects(course_b.id, a.filed[0].artifact_id) == []
    from app.main import app

    with TestClient(app) as client:
        response = client.get(
            f"/api/courses/{course_b.id}/artifacts/{a.filed[0].artifact_id}/objects"
        )
    assert response.status_code == 404
    assert b.filed[0].artifact_id != a.filed[0].artifact_id


def test_failed_artifact_can_be_retried_without_duplicate_manifest_rows(store):
    course = store.create_course("Retry")
    bad = import_files(store, course.id, [("data.ipynb", b"not json")])
    assert bad.errors and store.load_artifacts(course.id)[0].extraction_status == "failed"
    good_bytes = _ipynb()
    good = import_files(store, course.id, [("data.ipynb", good_bytes)])
    assert good.errors == []
    assert len(store.load_artifacts(course.id)) == 1
    assert store.load_artifacts(course.id)[0].version == 2


def test_native_view_api_paginates_sheets_and_serves_reference_pdf_render(store):
    course = store.create_course("Native views")
    syllabus = (Path(__file__).parent / "fixtures" / "syllabus.pdf").read_bytes()
    result = import_files(store, course.id, [
        ("L1/syllabus.pdf", syllabus), ("L2/book.xlsx", _xlsx()),
    ])
    by_kind = {item.kind: item for item in result.filed}
    from app.main import app

    with TestClient(app) as client:
        workbook = client.get(
            f"/api/courses/{course.id}/artifacts/{by_kind['xlsx'].artifact_id}/view",
            params={"sheet": "Returns", "limit": 2},
        )
        assert workbook.status_code == 200
        body = workbook.json()
        assert body["selected_sheet"] == "Returns" and body["limit"] == 2
        assert body["total_objects"] > 2 and body["has_more"] is True
        pdf_objects = client.get(
            f"/api/courses/{course.id}/artifacts/{by_kind['pdf'].artifact_id}/objects"
        ).json()
        page = next(obj for obj in pdf_objects if obj["object_type"] == "page")
        render = client.get(
            f"/api/courses/{course.id}/artifacts/{by_kind['pdf'].artifact_id}/render/{page['locator']['page']}"
        )
        assert render.status_code == 200 and render.content.startswith(b"\x89PNG")


def test_embedded_docx_and_notebook_images_render_from_course_scoped_objects(store):
    course = store.create_course("Embedded visuals")
    result = import_files(store, course.id, [("visual.docx", _docx()), ("plot.ipynb", _ipynb())])
    by_kind = {item.kind: item for item in result.filed}
    objects = store.load_learning_objects(course.id)
    doc_image = next(obj for obj in objects if obj.artifact_id == by_kind["docx"].artifact_id and obj.object_type == "image")
    plot = next(obj for obj in objects if obj.artifact_id == by_kind["ipynb"].artifact_id and "image/png" in obj.data.get("mime_types", []))
    from app.main import app

    with TestClient(app) as client:
        for artifact_id, object_id in ((doc_image.artifact_id, doc_image.id), (plot.artifact_id, plot.id)):
            response = client.get(f"/api/courses/{course.id}/artifacts/{artifact_id}/objects/{object_id}/media")
            assert response.status_code == 200
            assert response.headers["content-type"] == "image/png"
            assert response.content.startswith(b"\x89PNG")
        foreign = store.create_course("Foreign visuals")
        rejected = client.get(f"/api/courses/{foreign.id}/artifacts/{plot.artifact_id}/objects/{plot.id}/media")
        assert rejected.status_code == 404


def test_malformed_embedded_docx_image_is_not_served_as_trusted_media(store):
    original = BytesIO(_docx())
    damaged = BytesIO()
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(damaged, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename.startswith("word/media/"):
                content = b"not an image"
            target.writestr(info, content)
    course = store.create_course("Damaged media")
    imported = import_files(store, course.id, [("damaged.docx", damaged.getvalue())])
    artifact_id = imported.filed[0].artifact_id
    image = next(obj for obj in store.load_learning_objects(course.id, artifact_id) if obj.object_type == "image")
    from app.main import app

    with TestClient(app) as client:
        response = client.get(f"/api/courses/{course.id}/artifacts/{artifact_id}/objects/{image.id}/media")
    assert response.status_code == 404 and "invalid" in response.json()["error"].lower()


def test_selected_context_question_is_grounded_cited_and_course_scoped(store, monkeypatch):
    monkeypatch.setenv("LC_FAKE_MODEL", "1")
    course_a = store.create_course("Question A")
    course_b = store.create_course("Question B")
    imported_a = import_files(store, course_a.id, [("facts.csv", b"term,value\nalpha,42\n")])
    imported_b = import_files(store, course_b.id, [("facts.csv", b"term,value\nbeta,7\n")])
    aid = imported_a.filed[0].artifact_id
    obj = next(item for item in store.load_learning_objects(course_a.id, aid) if item.text == "value")
    foreign = imported_b.filed[0].artifact_id
    from app.main import app

    with TestClient(app) as client:
        response = client.post(f"/api/courses/{course_a.id}/questions", json={
            "question": "What does the selected field show?",
            "selections": [{"artifact_id": aid, "object_id": obj.id}],
        })
        assert response.status_code == 200
        body = response.json()
        assert "[S1]" in body["answer"]
        assert '"missing_count": 0' in body["answer"] and '"row_count": 1' in body["answer"]
        assert body["citations"][0]["object_id"] == obj.id
        assert body["citations"][0]["locator"]["dataset_field"] == "value"
        assert body["remote_disclosure"]["sent_remote"] is False
        rejected = client.post(f"/api/courses/{course_a.id}/questions", json={
            "question": "Leak?",
            "selections": [{"artifact_id": foreign, "object_id": "anything"}],
        })
        assert rejected.status_code == 404


def test_frontend_exposes_native_view_and_selected_question_controls():
    html = (Path(__file__).parents[1] / "web" / "index.html").read_text("utf-8")
    script = (Path(__file__).parents[1] / "web" / "app.js").read_text("utf-8")
    for marker in ("data-native-viewer", "source-location", "copy-source-link", "ask-selected-context"):
        assert marker in html
    assert "/questions" in script and "/artifacts/" in script
    assert "sheet-metadata" in script and "native-embedded" in script
