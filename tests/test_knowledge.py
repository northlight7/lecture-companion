"""Course-scoped typed search, concept graph, and staleness regression tests."""

from __future__ import annotations

import json
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.importer import import_files
from app.knowledge import build_index, concept_graph, index_status, search


def _docx(text: str) -> bytes:
    out = BytesIO()
    doc = Document()
    doc.add_heading("Regression foundations", 1)
    doc.add_paragraph(text)
    doc.save(out)
    return out.getvalue()


def _xlsx() -> bytes:
    out = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ridge penalty"
    sheet.append(["lambda", "coefficient"])
    sheet.append([0.1, "=A2*2"])
    workbook.save(out)
    return out.getvalue()


def _notebook(signal: str) -> bytes:
    return json.dumps({
        "cells": [{
            "cell_type": "markdown", "id": "regularization", "metadata": {},
            "source": [f"Ridge regularization shrinks coefficients. {signal}"],
        }],
        "metadata": {}, "nbformat": 4, "nbformat_minor": 5,
    }).encode()


def test_local_search_crosses_artifact_types_with_exact_typed_links(store):
    course = store.create_course("Cross artifact analytics")
    imported = import_files(store, course.id, [
        ("lecture.docx", _docx("Linear regression uses least squares before ridge regularization.")),
        ("practice.xlsx", _xlsx()),
        ("training.csv", b"feature,target\n1,2\n"),
        ("analysis.ipynb", _notebook("COURSE_A_ONLY pd.read_csv('training.csv')")),
    ])
    assert imported.errors == []

    result = search(store, course.id, "shrinkage", limit=20)
    assert result["rebuilt"] is True
    assert result["expanded_concepts"] == ["regularization"]
    kinds = {row["artifact_kind"] for row in result["results"]}
    assert {"docx", "xlsx", "ipynb", "csv"} <= kinds
    assert all(row["artifact_id"] and row["object_id"] and row["locator"] for row in result["results"])
    assert any(row["matched_by"] == "concept alias" for row in result["results"])
    related_dataset = next(row for row in result["results"] if row["matched_by"] == "artifact relationship")
    assert related_dataset["artifact_kind"] == "csv" and "row_count" in related_dataset["snippet"]


def test_concept_graph_links_exact_evidence_and_labels_prerequisites_as_candidates(store):
    course = store.create_course("Concept relationships")
    import_files(store, course.id, [
        ("lecture.docx", _docx("Regression residuals motivate ridge regularization.")),
        ("observations.csv", b"x,y\n1,2\n"),
        ("lab.ipynb", _notebook("model evaluation uses RMSE and pd.read_csv('observations.csv')")),
        ("penalties.xlsx", _xlsx()),
    ])
    graph = concept_graph(store, course.id, limit=30)
    concepts = {node["label"]: node for node in graph["nodes"] if node["node_type"] == "concept"}
    assert "regression" in concepts and "regularization" in concepts
    assert concepts["regularization"]["artifact_count"] >= 3
    appears = [edge for edge in graph["edges"] if edge["relation"] == "appears_in"]
    assert appears and all(edge["object_ids"] and "exact source" in edge["reason"] for edge in appears)
    prerequisite = next(edge for edge in graph["edges"] if edge["relation"] == "prerequisite_candidate")
    assert prerequisite["certainty"] == "modeled candidate"
    explicit = next(edge for edge in graph["edges"] if edge["relation"] == "explicit_reference")
    assert explicit["source_object_id"] and explicit["target_object_id"]
    assert explicit["certainty"] == "measured filename reference"
    assert "not source facts" in graph["warning"]


def test_changed_source_marks_index_stale_and_rebuild_removes_superseded_content(store):
    course = store.create_course("Versioned search")
    import_files(store, course.id, [("notes.docx", _docx("OBSOLETE_MARKER regression"))])
    build_index(store, course.id)
    assert index_status(store, course.id)["stale"] is False

    import_files(store, course.id, [("notes.docx", _docx("CURRENT_MARKER regression"))])
    stale = index_status(store, course.id)
    assert stale["stale"] is True and "source versions changed" in stale["reasons"]
    current = search(store, course.id, "CURRENT_MARKER")
    assert current["rebuilt"] is True and current["results"]
    assert search(store, course.id, "OBSOLETE_MARKER")["results"] == []


def test_search_and_graph_never_cross_course_boundaries(store):
    course_a = store.create_course("Analytics A")
    course_b = store.create_course("Analytics B")
    import_files(store, course_a.id, [("a.ipynb", _notebook("A_SECRET_SIGNAL"))])
    import_files(store, course_b.id, [("b.ipynb", _notebook("B_SECRET_SIGNAL"))])
    a_blob = json.dumps(search(store, course_a.id, "regression")) + json.dumps(concept_graph(store, course_a.id))
    b_blob = json.dumps(search(store, course_b.id, "regression")) + json.dumps(concept_graph(store, course_b.id))
    assert "B_SECRET_SIGNAL" not in a_blob
    assert "A_SECRET_SIGNAL" not in b_blob


def test_knowledge_api_is_course_scoped_and_reports_staleness(store):
    course = store.create_course("Knowledge API")
    imported = import_files(store, course.id, [("lab.ipynb", _notebook("API_SIGNAL"))])
    foreign = store.create_course("Foreign")
    from app.main import app

    with TestClient(app) as client:
        before = client.get(f"/api/courses/{course.id}/knowledge/status")
        assert before.status_code == 200 and before.json()["stale"] is True
        found = client.get(f"/api/courses/{course.id}/search", params={"q": "regularization"})
        assert found.status_code == 200 and found.json()["results"]
        graph = client.get(f"/api/courses/{course.id}/concepts")
        assert graph.status_code == 200 and graph.json()["nodes"]
        absent = client.get(f"/api/courses/{foreign.id}/artifacts/{imported.filed[0].artifact_id}/objects")
        assert absent.status_code == 404
        missing = client.get("/api/courses/no-such-course/search", params={"q": "signal"})
        assert missing.status_code == 404
