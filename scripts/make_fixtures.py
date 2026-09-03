#!/usr/bin/env python
"""Generate the committed test fixtures in tests/fixtures/.

Deterministic and idempotent: fixed dates, no randomness, same bytes-worth of
content on every run. Run it with:

    .venv/bin/python scripts/make_fixtures.py
"""

from __future__ import annotations

import datetime as _dt
import sys
from pathlib import Path

from reportlab import rl_config

rl_config.invariant = 1  # fixed /CreationDate and document id -> reproducible PDFs

from reportlab.lib.pagesizes import landscape, A4  # noqa: E402
from reportlab.pdfgen import canvas as _canvas  # noqa: E402

REPO_ROOT = Path(__file__).resolve().parent.parent
FIXTURES = REPO_ROOT / "tests" / "fixtures"

FIXED_DATE = _dt.datetime(2026, 1, 1, 0, 0, 0)

# --------------------------------------------------------------------------
# content: one coherent topic, so grounding tests have something real to check
# --------------------------------------------------------------------------

LECTURE_W1: list[tuple[str, list[str]]] = [
    ("Introduction to Databases", [
        "COMP5211 - Week 1",
        "What a database management system does",
        "Why files on disk are not enough",
        "Roadmap: model, tables, keys, SQL, joins",
    ]),
    ("The Relational Model", [
        "Data is stored as relations, not pointers",
        "A relation is a set of tuples",
        "Every attribute has a domain",
        "Order of rows carries no meaning",
        "Proposed by E. F. Codd in 1970",
    ]),
    ("Tables, Rows and Columns", [
        "A table is one relation on disk",
        "A row is one tuple, one real-world fact",
        "A column is one attribute with one type",
        "NULL means unknown, not zero",
    ]),
    ("Keys and Constraints", [
        "A primary key identifies a row uniquely",
        "A foreign key points at another table's key",
        "Referential integrity forbids dangling references",
        "UNIQUE and NOT NULL are checked on write",
    ]),
    ("SQL SELECT", [
        "SELECT columns FROM table WHERE condition",
        "WHERE filters rows before grouping",
        "ORDER BY sorts the result, not the table",
        "LIMIT caps how many rows come back",
    ]),
    ("Joins and Normalisation", [
        "An INNER JOIN keeps only matching rows",
        "A LEFT JOIN keeps unmatched left rows too",
        "First normal form: no repeating groups",
        "Third normal form removes transitive dependencies",
        "Normalise to stop update anomalies",
    ]),
]

SYLLABUS_PAGES: list[tuple[str, list[str]]] = [
    ("COMP5211 Introduction to Database Systems - Course Syllabus", [
        "Course description. This course introduces the design, implementation and use of "
        "relational database systems for students entering the MSc programme with no prior "
        "database background. It begins with the relational model as a mathematical object, "
        "moves through schema design and normalisation, and then covers the SQL language in "
        "enough depth that students can express realistic analytical queries without help. "
        "The second half of the term turns to what happens underneath the query: storage "
        "layout, indexing structures, transaction management and the guarantees a database "
        "makes when several users write at the same time. The course is deliberately "
        "practical. Every concept introduced in lecture is exercised the same week against a "
        "live PostgreSQL instance provided to each student, because experience shows that "
        "students who only read about isolation levels do not retain them.",
        "Learning outcomes. On successful completion of this course a student will be able "
        "to: design a normalised relational schema from an informal description of a "
        "business domain and justify the normal form reached; write correct SQL queries "
        "involving joins, aggregation, subqueries and window functions; read an execution "
        "plan and explain why the planner chose a particular access path; reason about the "
        "correctness of concurrent transactions using the ACID properties and the standard "
        "isolation levels; and evaluate the trade-offs between relational and non-relational "
        "storage for a stated workload. These outcomes are assessed cumulatively rather than "
        "in isolation, so the final examination assumes everything earlier in the term.",
        "Prerequisite. Students are expected to be comfortable with elementary discrete "
        "mathematics, in particular sets, relations and functions, and to be able to write "
        "and debug a small program in any imperative language. No prior exposure to SQL or "
        "to any particular database product is assumed or required. Students who have not "
        "programmed for several years should contact the instructor during the first week.",
    ]),
    ("Assessment, Policies and Resources", [
        "Grading. The final grade is composed of four written assignments worth forty per "
        "cent in total, a midterm examination worth twenty per cent, a term project worth "
        "fifteen per cent, and a final examination worth twenty-five per cent. Assignments "
        "are released on Monday and are due at the end of the following week. Late work "
        "loses ten per cent of its available marks per calendar day and is not accepted more "
        "than five days after the deadline, except where documented medical or compassionate "
        "grounds apply. The term project is completed in pairs and is assessed on the "
        "quality of the schema design and the clarity of the written justification, not on "
        "the volume of code produced.",
        "Office hours. The instructor holds office hours on Tuesday and Thursday afternoons "
        "in Room 3521 and answers questions on the course forum within one working day. "
        "Teaching assistants run an additional consultation session before each assignment "
        "deadline. Students are encouraged to bring partially working queries rather than "
        "waiting until they have a complete solution, since most of the value of a "
        "consultation lies in the diagnosis rather than the answer.",
        "Textbook. The recommended textbook is Database System Concepts by Silberschatz, "
        "Korth and Sudarshan. The PostgreSQL manual is treated as required reading for the "
        "chapters on indexing and transactions, and specific sections are listed on the "
        "weekly schedule. No purchase is necessary: a copy of the textbook is held on "
        "reserve in the university library and the manual is freely available online.",
        "Academic honesty. All submitted work must be the student's own. Discussion of "
        "concepts between students is encouraged and expected, but the sharing of written "
        "answers, query text or project code is not permitted and is treated as plagiarism "
        "under the university's academic integrity regulations. Use of a generative language "
        "model to produce submitted answers must be disclosed in the submission. Suspected "
        "cases are referred to the departmental committee without exception, and the "
        "penalties range from a zero on the assessment to failure of the course.",
    ]),
]

PPTX_W2: list[tuple[str, list[str], str]] = [
    ("Indexes", [
        "A B-tree index turns a scan into a lookup",
        "Indexes speed reads and slow writes",
        "The planner may ignore an index it thinks is useless",
    ], ""),
    ("Transactions", [
        "A transaction is an all-or-nothing unit of work",
        "BEGIN, then COMMIT or ROLLBACK",
        "Partial effects are never visible to others",
    ], "Reminder for the lecturer: link this back to the referential integrity example "
       "from week one, where a half-applied update left a dangling foreign key. That is "
       "exactly the failure a transaction prevents."),
    ("ACID", [
        "Atomicity: all of it or none of it",
        "Consistency: constraints hold before and after",
        "Isolation: concurrent transactions do not interleave visibly",
        "Durability: a committed write survives a crash",
    ], ""),
    ("Query Plans", [
        "EXPLAIN shows the plan the optimiser chose",
        "Sequential scan versus index scan",
        "Row-count estimates drive join order",
    ], ""),
]


# --------------------------------------------------------------------------


def _slide_pdf(dest: Path, pages: list[tuple[str, list[str]]], title: str) -> int:
    size = landscape(A4)
    c = _canvas.Canvas(str(dest), pagesize=size, invariant=1)
    c.setTitle(title)
    c.setAuthor("Lecture Companion fixtures")
    c.setSubject("Introduction to Databases")
    w, h = size
    for heading, bullets in pages:
        c.setFont("Helvetica-Bold", 30)
        c.drawString(60, h - 90, heading)
        c.setFont("Helvetica", 20)
        y = h - 160
        for b in bullets:
            c.drawString(80, y, "-  " + b)
            y -= 42
        c.showPage()
    c.save()
    return len(pages)


def _prose_pdf(dest: Path, pages: list[tuple[str, list[str]]], title: str) -> int:
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph

    size = A4
    c = _canvas.Canvas(str(dest), pagesize=size, invariant=1)
    c.setTitle(title)
    c.setAuthor("Lecture Companion fixtures")
    w, h = size
    body = ParagraphStyle("body", fontName="Times-Roman", fontSize=10.5, leading=13.5,
                          spaceAfter=8)
    for heading, paras in pages:
        c.setFont("Times-Bold", 14)
        c.drawString(20 * mm, h - 22 * mm, heading)
        y = h - 32 * mm
        for text in paras:
            p = Paragraph(text, body)
            avail_w = w - 40 * mm
            _pw, ph = p.wrap(avail_w, y - 20 * mm)
            p.drawOn(c, 20 * mm, y - ph)
            y -= ph + 8
        c.showPage()
    c.save()
    return len(pages)


def _pptx(dest: Path) -> int:
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    blank = prs.slide_layouts[1]        # Title and Content
    for heading, bullets, notes in PPTX_W2:
        slide = prs.slides.add_slide(blank)
        slide.shapes.title.text = heading
        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for b in bullets[1:]:
            para = body.add_paragraph()
            para.text = b
            para.level = 0
        for para in body.paragraphs:
            for run in para.runs:
                run.font.size = Pt(20)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    cp = prs.core_properties
    cp.title = "COMP5211 Week 2 - Indexes and Transactions"
    cp.author = "Lecture Companion fixtures"
    cp.created = FIXED_DATE
    cp.modified = FIXED_DATE
    cp.revision = 1
    prs.save(str(dest))
    return len(PPTX_W2)


def make_fixtures(out_dir: Path | None = None) -> dict[str, int]:
    out = Path(out_dir) if out_dir else FIXTURES
    out.mkdir(parents=True, exist_ok=True)
    counts = {
        "lecture_w1.pdf": _slide_pdf(
            out / "lecture_w1.pdf", LECTURE_W1, "COMP5211 Week 1 - Introduction to Databases"
        ),
        "syllabus.pdf": _prose_pdf(
            out / "syllabus.pdf", SYLLABUS_PAGES, "COMP5211 Course Syllabus"
        ),
        "lecture_w2.pptx": _pptx(out / "lecture_w2.pptx"),
    }
    return counts


def main(argv: list[str]) -> int:
    out = Path(argv[1]) if len(argv) > 1 else FIXTURES
    counts = make_fixtures(out)
    for name, n in counts.items():
        path = out / name
        unit = "slides" if path.suffix == ".pptx" else "pages"
        print(f"wrote {path}  ({n} {unit}, {path.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
