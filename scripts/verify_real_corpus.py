#!/usr/bin/env python3
"""Import the read-only five-course corpus and report structural fidelity metrics."""

from __future__ import annotations

import argparse
import hashlib
import json
import tempfile
import warnings
from collections import Counter
from pathlib import Path

from app.course import CourseStore
from app.importer import import_files


def _structure_failures(path: Path, kind: str, objects) -> list[str]:
    actual = Counter(obj.object_type for obj in objects)
    failures: list[str] = []
    if kind == "pdf":
        import pypdfium2 as pdfium

        doc = pdfium.PdfDocument(str(path))
        expected = len(doc)
        doc.close()
        if actual["page"] != expected:
            failures.append(f"page objects {actual['page']} != source pages {expected}")
    elif kind == "pptx":
        from pptx import Presentation

        deck = Presentation(str(path))
        expected_notes = 0
        for slide in deck.slides:
            try:
                frame = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
                expected_notes += int(bool((frame.text or "").strip())) if frame is not None else 0
            except Exception:
                pass
        if actual["slide"] != len(deck.slides) or actual["speaker_note"] != expected_notes:
            failures.append("slide or speaker-note count differs from source")
    elif kind == "docx":
        from docx import Document

        doc = Document(str(path))
        expected_blocks = 0
        expected_tables = 0
        for item in doc.iter_inner_content():
            if hasattr(item, "rows"):
                expected_tables += 1
                expected_blocks += 1
            elif (getattr(item, "text", "") or "").strip():
                expected_blocks += 1
        expected_images = sum("image" in rel.reltype for rel in doc.part.rels.values())
        expected_links = sum(bool("hyperlink" in rel.reltype and getattr(rel, "target_ref", "")) for rel in doc.part.rels.values())
        if actual["document_block"] + actual["prompt"] + actual["table"] != expected_blocks:
            failures.append("ordered Word block count differs from source")
        if actual["table"] != expected_tables or actual["image"] != expected_images or actual["hyperlink"] != expected_links:
            failures.append("Word table, image, or hyperlink count differs from source")
    elif kind == "xlsx":
        import openpyxl

        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            workbook = openpyxl.load_workbook(path, data_only=False, read_only=False)
        expected_cells = sum(
            cell.value is not None for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row
        )
        expected_formulas = sum(
            cell.data_type == "f" for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row
        )
        expected_charts = sum(len(sheet._charts) for sheet in workbook.worksheets)
        expected_tables = sum(len(sheet.tables) for sheet in workbook.worksheets)
        expected_features = {
            "hidden_sheets": sum(sheet.sheet_state != "visible" for sheet in workbook.worksheets),
            "merged_ranges": sum(len(sheet.merged_cells.ranges) for sheet in workbook.worksheets),
            "hidden_rows": sum(sum(bool(dim.hidden) for dim in sheet.row_dimensions.values()) for sheet in workbook.worksheets),
            "hidden_columns": sum(sum(bool(dim.hidden) for dim in sheet.column_dimensions.values()) for sheet in workbook.worksheets),
            "filters": sum(bool(sheet.auto_filter.ref) for sheet in workbook.worksheets),
            "conditional_formatting": sum(len(sheet.conditional_formatting) for sheet in workbook.worksheets),
            "validations": sum(len(sheet.data_validations.dataValidation) for sheet in workbook.worksheets),
        }
        sheet_objects = [obj for obj in objects if obj.object_type == "sheet"]
        actual_features = {
            "hidden_sheets": sum(obj.data.get("state") != "visible" for obj in sheet_objects),
            "merged_ranges": sum(len(obj.data.get("merged_ranges", [])) for obj in sheet_objects),
            "hidden_rows": sum(len(obj.data.get("hidden_rows", [])) for obj in sheet_objects),
            "hidden_columns": sum(len(obj.data.get("hidden_columns", [])) for obj in sheet_objects),
            "filters": sum(bool(obj.data.get("auto_filter")) for obj in sheet_objects),
            "conditional_formatting": sum(int(obj.data.get("conditional_formatting_rules", 0)) for obj in sheet_objects),
            "validations": sum(len(obj.data.get("data_validations", [])) for obj in sheet_objects),
        }
        formulas = sum(obj.data.get("formula") is not None for obj in objects if obj.object_type == "cell")
        if actual["sheet"] != len(workbook.worksheets) or actual["cell"] != expected_cells:
            failures.append("workbook sheet or non-empty cell count differs from source")
        if formulas != expected_formulas or actual["chart"] != expected_charts or actual["range"] < expected_tables:
            failures.append("workbook formula, chart, or table count differs from source")
        if actual_features != expected_features:
            failures.append(f"workbook feature inventory differs: {actual_features} != {expected_features}")
    elif kind == "csv":
        import csv
        from io import StringIO

        text = path.read_text("utf-8-sig")
        rows = list(csv.DictReader(StringIO(text)))
        fields = list(rows[0].keys()) if rows else next(csv.reader(StringIO(text)), [])
        dataset = next(obj for obj in objects if obj.object_type == "dataset")
        if dataset.data.get("row_count") != len(rows) or actual["dataset_field"] != len(fields):
            failures.append("dataset row or field count differs from source")
    elif kind == "ipynb":
        notebook = json.loads(path.read_text("utf-8"))
        expected_outputs = sum(len(cell.get("outputs", []) or []) for cell in notebook["cells"])
        if actual["notebook_cell"] != len(notebook["cells"]) or actual["notebook_output"] != expected_outputs:
            failures.append("notebook cell or output count differs from source")
    return failures


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=Path("/Users/north/Desktop/Courses"))
    parser.add_argument("--scratch-parent", type=Path, default=Path(__file__).parents[1] / ".internal")
    args = parser.parse_args()
    source = args.source.resolve()
    scratch_parent = args.scratch_parent.resolve()
    scratch_parent.mkdir(parents=True, exist_ok=True)
    files = sorted(path for path in source.rglob("*") if path.is_file() and path.name != ".DS_Store")
    failures: list[str] = []
    kind_counts: Counter[str] = Counter()
    object_counts: Counter[str] = Counter()
    degraded_artifacts = 0
    with tempfile.TemporaryDirectory(prefix="real-corpus-", dir=scratch_parent) as tmp:
        store = CourseStore(Path(tmp) / "Courses")
        for course_dir in sorted(path for path in source.iterdir() if path.is_dir()):
            course = store.create_course(course_dir.name)
            payload = [
                (path.relative_to(course_dir).as_posix(), path.read_bytes())
                for path in sorted(course_dir.rglob("*"))
                if path.is_file() and path.name != ".DS_Store"
            ]
            result = import_files(store, course.id, payload)
            failures.extend(result.errors)
            for artifact in store.load_artifacts(course.id):
                kind_counts[artifact.kind] += 1
                original = store.course_dir(course.id) / artifact.stored_path
                if hashlib.sha256(original.read_bytes()).hexdigest() != artifact.content_hash:
                    failures.append(f"{course.title}/{artifact.source_path}: original hash mismatch")
                objects = store.load_learning_objects(course.id, artifact.id)
                if len(objects) != artifact.object_count or not objects:
                    failures.append(f"{course.title}/{artifact.source_path}: object count mismatch")
                for obj in objects:
                    object_counts[obj.object_type] += 1
                    if obj.locator.artifact_id != artifact.id:
                        failures.append(f"{course.title}/{artifact.source_path}: locator mismatch")
                source_file = course_dir / artifact.source_path
                for detail in _structure_failures(source_file, artifact.kind, objects):
                    failures.append(f"{course.title}/{artifact.source_path}: {detail}")
                if artifact.extraction_warnings:
                    degraded_artifacts += 1
                    if artifact.extraction_quality >= 1.0:
                        failures.append(f"{course.title}/{artifact.source_path}: warning not reflected in quality")
                    if artifact.kind == "xlsx" and not any(obj.object_type == "package_part" for obj in objects):
                        failures.append(f"{course.title}/{artifact.source_path}: unsupported package part not inventoried")
        report = {
            "source_files": len(files),
            "artifacts": sum(kind_counts.values()),
            "formats": dict(sorted(kind_counts.items())),
            "learning_objects": sum(object_counts.values()),
            "object_types": dict(sorted(object_counts.items())),
            "degraded_artifacts": degraded_artifacts,
            "failures": failures,
        }
        print(json.dumps(report, indent=2, sort_keys=True))
        expected = {"pdf": 8, "pptx": 3, "docx": 2, "xlsx": 3, "csv": 11, "ipynb": 2}
        return 0 if len(files) == 29 and dict(kind_counts) == expected and not failures else 1


if __name__ == "__main__":
    raise SystemExit(main())
